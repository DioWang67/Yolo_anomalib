from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.services.inspection_excel_report import (
    InspectionExcelReportCancelled,
    InspectionExcelReportService,
)
from core.services.inspection_history import InspectionHistoryFilters
from core.services.inspection_repository import InspectionRepository


def _insert(
    repository: InspectionRepository,
    root: Path,
    *,
    name: str,
    timestamp: str,
    status: str,
    product: str = "Cable1",
    station: str = "A",
    reasons: object = None,
) -> None:
    repository.upsert_snapshot(
        {
            "timestamp": timestamp,
            "status": status,
            "detector": "yolo",
            "product": product,
            "equipment": {
                "station": station,
                "machine_id": "M-01",
                "work_order": "WO-007",
                "camera_id": "CAM-A",
            },
            "model_info": {"model_version": "v3"},
            "inference_time": 0.125,
            "fail_reasons": reasons or [],
            "artifacts": {
                "original_path": str(root / f"{name}-original.jpg"),
                "annotated_path": str(root / f"{name}-annotated.jpg"),
            },
        },
        snapshot_path=root / f"{name}.json",
    )


def test_excel_report_exports_all_filtered_rows_and_summary(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "Result" / "inspection_records.sqlite3"
    repository = InspectionRepository(database_path)
    _insert(
        repository,
        tmp_path,
        name="pass",
        timestamp="2026-07-28T08:00:00+00:00",
        status="PASS",
    )
    _insert(
        repository,
        tmp_path,
        name="fail-one",
        timestamp="2026-07-28T09:00:00+00:00",
        status="DETECTION_FAIL",
        reasons=[{"code": "POSITION_SHIFT"}],
    )
    _insert(
        repository,
        tmp_path,
        name="fail-two",
        timestamp="2026-07-28T10:00:00+00:00",
        status="NG",
        reasons=["MISSING"],
    )
    _insert(
        repository,
        tmp_path,
        name="other-station",
        timestamp="2026-07-28T11:00:00+00:00",
        status="NG",
        station="B",
    )
    destination = tmp_path / "filtered.xlsx"

    result = InspectionExcelReportService(database_path).export(
        destination,
        InspectionHistoryFilters(
            product="Cable1",
            station="A",
            status="FAIL",
            limit=1,
        ),
        language="zh",
    )

    assert result.path == destination.resolve()
    assert result.record_count == 2
    workbook = load_workbook(destination, read_only=False, data_only=True)
    assert workbook.sheetnames == ["報表摘要", "檢測明細", "異常統計"]

    summary = dict(
        row
        for row in workbook["報表摘要"].iter_rows(
            min_row=3,
            max_col=2,
            values_only=True,
        )
        if row[0]
    )
    assert summary["產品"] == "Cable1"
    assert summary["工位"] == "A"
    assert summary["狀態"] == "FAIL"
    assert summary["符合筆數"] == 2
    assert summary["PASS"] == 0
    assert summary["NG"] == 2
    assert summary["良率"] == 0

    details = workbook["檢測明細"]
    assert details.max_row == 3
    assert details["B2"].value == "2026-07-28 18:00:00"
    assert details["C2"].value == "NG"
    assert details["L2"].value == "MISSING"
    assert details["O2"].hyperlink is not None
    assert details.auto_filter.ref == "A1:R3"

    reasons = {
        row[0]: row[1]
        for row in workbook["異常統計"].iter_rows(
            min_row=2,
            max_col=2,
            values_only=True,
        )
    }
    assert reasons == {"MISSING": 1, "POSITION_SHIFT": 1}
    assert list(tmp_path.glob(".*.tmp.xlsx")) == []


def test_excel_report_creates_empty_workbook_without_database(
    tmp_path: Path,
) -> None:
    destination = tmp_path / "empty.xlsx"

    result = InspectionExcelReportService(
        tmp_path / "missing.sqlite3"
    ).export(destination, InspectionHistoryFilters(), language="en")

    assert result.record_count == 0
    workbook = load_workbook(destination, read_only=False)
    assert workbook.sheetnames == [
        "Summary",
        "Inspection details",
        "Failure reasons",
    ]
    assert workbook["Inspection details"].max_row == 1


@pytest.mark.parametrize(
    "destination",
    (
        "report.csv",
        "missing/report.xlsx",
    ),
)
def test_excel_report_rejects_invalid_destination(
    tmp_path: Path,
    destination: str,
) -> None:
    with pytest.raises(ValueError):
        InspectionExcelReportService(
            tmp_path / "missing.sqlite3"
        ).export(
            tmp_path / destination,
            InspectionHistoryFilters(),
        )


def test_excel_report_escapes_untrusted_formula_text(tmp_path: Path) -> None:
    database_path = tmp_path / "inspection_records.sqlite3"
    repository = InspectionRepository(database_path)
    _insert(
        repository,
        tmp_path,
        name="formula",
        timestamp="2026-07-28T08:00:00+00:00",
        status="NG",
        product="=2+2",
        reasons=["@SUM(A1:A2)"],
    )
    destination = tmp_path / "safe.xlsx"

    InspectionExcelReportService(database_path).export(
        destination,
        InspectionHistoryFilters(product="=2+2"),
    )

    workbook = load_workbook(destination, data_only=False)
    assert workbook["報表摘要"]["B4"].value == "'=2+2"
    assert workbook["檢測明細"]["D2"].value == "'=2+2"
    assert workbook["檢測明細"]["L2"].value == "'@SUM(A1:A2)"


def test_excel_report_cancellation_leaves_no_partial_file(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "inspection_records.sqlite3"
    repository = InspectionRepository(database_path)
    for index in range(3):
        _insert(
            repository,
            tmp_path,
            name=f"cancel-{index}",
            timestamp=f"2026-07-28T08:00:0{index}+00:00",
            status="PASS",
        )
    destination = tmp_path / "cancelled.xlsx"
    progress: list[tuple[int, int]] = []

    with pytest.raises(InspectionExcelReportCancelled):
        InspectionExcelReportService(database_path).export(
            destination,
            InspectionHistoryFilters(),
            progress_callback=lambda current, total: progress.append(
                (current, total)
            ),
            is_cancelled=lambda: bool(progress),
        )

    assert progress == [(0, 3)]
    assert destination.exists() is False
    assert list(tmp_path.glob(".*.tmp.xlsx")) == []
