from __future__ import annotations

import threading

from openpyxl import load_workbook

from core.services.results import excel_buffer as excel_buffer_module
from core.services.results.excel_buffer import ExcelWorkbookBuffer


class DummyLogger:
    def __init__(self) -> None:
        self.warnings: list[str] = []
        self.infos: list[str] = []

    def warning(self, message, *args):
        self.warnings.append(message % args if args else message)

    def info(self, message, *args):
        self.infos.append(message % args if args else message)

    def error(self, message, *args):
        pass


def test_excel_buffer_rebuilds_corrupt_workbook(tmp_path):
    path = tmp_path / "results.xlsx"
    path.write_text("not a real xlsx", encoding="utf-8")
    logger = DummyLogger()

    buffer = ExcelWorkbookBuffer(
        path=str(path),
        columns=["id", "status"],
        buffer_limit=10,
        logger=logger,
    )

    assert path.exists()
    assert buffer.ws.max_row == 1
    assert list(tmp_path.glob("results.xlsx.corrupt_*"))
    assert logger.warnings
    buffer.close()


def test_failed_flush_keeps_rows_buffered(monkeypatch, tmp_path):
    path = tmp_path / "results.xlsx"
    logger = DummyLogger()
    buffer = ExcelWorkbookBuffer(
        path=str(path),
        columns=["id", "status"],
        buffer_limit=10,
        logger=logger,
    )
    monkeypatch.setattr(excel_buffer_module.time, "sleep", lambda *_a: None)
    monkeypatch.setattr(
        excel_buffer_module.os,
        "replace",
        lambda *_a, **_k: (_ for _ in ()).throw(PermissionError("workbook open")),
    )

    buffer.append([1, "PASS"])
    result = buffer.flush()

    assert result is not None and result.success is False
    assert buffer.pending_rows() == 1
    assert buffer.ws.max_row == 1
    buffer.wb.close()


def test_transient_retry_does_not_duplicate_rows(monkeypatch, tmp_path):
    path = tmp_path / "results.xlsx"
    logger = DummyLogger()
    buffer = ExcelWorkbookBuffer(
        path=str(path),
        columns=["id", "status"],
        buffer_limit=10,
        logger=logger,
    )
    monkeypatch.setattr(excel_buffer_module.time, "sleep", lambda *_a: None)
    real_replace = excel_buffer_module.os.replace
    calls = 0

    def replace_once(source, destination):
        nonlocal calls
        calls += 1
        if calls == 1:
            raise PermissionError("temporary lock")
        return real_replace(source, destination)

    monkeypatch.setattr(excel_buffer_module.os, "replace", replace_once)

    buffer.append([1, "PASS"])
    result = buffer.flush()
    buffer.close()

    assert result is not None and result.success is True
    workbook = load_workbook(path)
    assert workbook.active.max_row == 2
    assert workbook.active.cell(row=2, column=1).value == 1
    workbook.close()


def test_async_flush_does_not_block_new_rows(monkeypatch, tmp_path):
    path = tmp_path / "results.xlsx"
    logger = DummyLogger()
    buffer = ExcelWorkbookBuffer(
        path=str(path),
        columns=["id", "status"],
        buffer_limit=10,
        logger=logger,
    )
    write_started = threading.Event()
    allow_write = threading.Event()
    real_write = buffer._write_rows

    def delayed_write(rows):
        write_started.set()
        assert allow_write.wait(timeout=2.0)
        return real_write(rows)

    monkeypatch.setattr(buffer, "_write_rows", delayed_write)
    buffer.append([1, "PASS"])
    buffer.flush_async()
    assert write_started.wait(timeout=2.0)

    # The slow workbook write owns only the write lock; producers can continue.
    buffer.append([2, "PASS"])
    assert buffer.next_test_id() == 3
    assert buffer.pending_rows() == 2

    allow_write.set()
    buffer.close()

    workbook = load_workbook(path)
    assert workbook.active.max_row == 3
    assert workbook.active.cell(row=2, column=1).value == 1
    assert workbook.active.cell(row=3, column=1).value == 2
    workbook.close()
