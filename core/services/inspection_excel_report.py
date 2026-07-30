"""Filtered Excel reports built from the inspection history database."""

from __future__ import annotations

import json
import os
import sqlite3
import uuid
from collections import Counter
from collections.abc import Callable
from contextlib import closing
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill

from core.services.inspection_history import (
    HISTORY_TIMEZONE,
    InspectionHistoryFilters,
    build_record_where_clause,
    parse_inspection_timestamp,
)

_MAX_EXPORT_RECORDS = 250_000
_DETAIL_COLUMNS = (
    "檢測編號",
    "檢測時間",
    "狀態",
    "產品",
    "工位",
    "檢測器",
    "推論耗時 (ms)",
    "機台",
    "工單",
    "相機",
    "模型版本",
    "異常原因",
    "複核結果",
    "異常分類",
    "原始影像",
    "標註影像",
    "熱圖",
    "結果快照",
)
_DETAIL_COLUMNS_EN = (
    "Inspection ID",
    "Inspection time",
    "Status",
    "Product",
    "Station",
    "Detector",
    "Inference (ms)",
    "Machine",
    "Work order",
    "Camera",
    "Model version",
    "Failure reasons",
    "Review outcome",
    "Failure category",
    "Original image",
    "Annotated image",
    "Heatmap",
    "Result snapshot",
)


class InspectionExcelReportError(RuntimeError):
    """A filtered inspection report could not be produced safely."""


class InspectionExcelReportCancelled(InspectionExcelReportError):
    """The operator cancelled an in-progress report export."""


@dataclass(frozen=True)
class InspectionExcelReport:
    path: Path
    record_count: int


class InspectionExcelReportService:
    """Create one atomic, read-only Excel snapshot for selected filters."""

    def __init__(self, database_path: str | Path) -> None:
        self.database_path = Path(database_path)

    def export(
        self,
        destination: str | Path,
        filters: InspectionHistoryFilters,
        *,
        language: str = "zh",
        progress_callback: Callable[[int, int], None] | None = None,
        is_cancelled: Callable[[], bool] | None = None,
    ) -> InspectionExcelReport:
        target = Path(destination)
        if target.suffix.lower() != ".xlsx":
            raise ValueError("Inspection report destination must end in .xlsx.")
        if not target.parent.is_dir():
            raise ValueError(
                f"Inspection report folder does not exist: {target.parent}"
            )

        query_filters = filters.normalized()
        temporary = target.with_name(
            f".{target.stem}.{uuid.uuid4().hex}.tmp.xlsx"
        )
        workbook: Workbook | None = None
        try:
            workbook = Workbook(write_only=True)
            summary_sheet = workbook.create_sheet(
                "報表摘要" if language == "zh" else "Summary"
            )
            detail_sheet = workbook.create_sheet(
                "檢測明細" if language == "zh" else "Inspection details"
            )
            reason_sheet = workbook.create_sheet(
                "異常統計" if language == "zh" else "Failure reasons"
            )
            record_count, status_counts, reason_counts = self._write_snapshot(
                detail_sheet,
                query_filters,
                language=language,
                progress_callback=progress_callback,
                is_cancelled=is_cancelled,
            )
            _raise_if_cancelled(is_cancelled)
            self._write_summary(
                summary_sheet,
                query_filters,
                record_count,
                status_counts,
                language=language,
            )
            self._write_reason_summary(
                reason_sheet,
                record_count,
                reason_counts,
                language=language,
            )
            workbook.save(temporary)
            os.replace(temporary, target)
        except InspectionExcelReportError:
            _discard_workbook(workbook, temporary)
            raise
        except (OSError, sqlite3.Error, ValueError) as exc:
            _discard_workbook(workbook, temporary)
            if isinstance(exc, ValueError):
                raise
            raise InspectionExcelReportError(
                f"Unable to export inspection Excel report: {exc}"
            ) from exc
        return InspectionExcelReport(
            path=target.resolve(),
            record_count=record_count,
        )

    def _write_snapshot(
        self,
        worksheet,
        filters: InspectionHistoryFilters,
        *,
        language: str,
        progress_callback: Callable[[int, int], None] | None,
        is_cancelled: Callable[[], bool] | None,
    ) -> tuple[int, Counter[str], Counter[str]]:
        columns = _DETAIL_COLUMNS if language == "zh" else _DETAIL_COLUMNS_EN
        worksheet.append(_header_cells(worksheet, columns))
        worksheet.freeze_panes = "A2"
        _set_detail_widths(worksheet)
        if not self.database_path.is_file():
            worksheet.auto_filter.ref = "A1:R1"
            if progress_callback is not None:
                progress_callback(0, 0)
            return 0, Counter(), Counter()

        where_sql, parameters = build_record_where_clause(filters)
        status_counts: Counter[str] = Counter()
        reason_counts: Counter[str] = Counter()
        record_count = 0
        try:
            with closing(self._connect()) as connection:
                connection.execute("BEGIN")
                total_row = connection.execute(
                    f"""
                    SELECT COUNT(*) AS count
                    FROM inspections
                    {where_sql}
                    """,
                    parameters,
                ).fetchone()
                total_count = int(total_row["count"] if total_row else 0)
                if total_count > _MAX_EXPORT_RECORDS:
                    raise InspectionExcelReportError(
                        "Inspection report exceeds the safe export limit "
                        f"of {_MAX_EXPORT_RECORDS:,} records. Narrow the "
                        "date or target filters and retry."
                    )
                if progress_callback is not None:
                    progress_callback(0, total_count)
                _raise_if_cancelled(is_cancelled)
                cursor = connection.execute(
                    f"""
                    SELECT
                        inspection_id, timestamp, status, detector, product,
                        station, machine_id, work_order, camera_id,
                        model_version, inference_time, decision_reasons_json,
                        review_outcome, failure_category, original_path,
                        annotated_path, heatmap_path, snapshot_path
                    FROM inspections
                    {where_sql}
                    ORDER BY timestamp DESC, inspection_id DESC
                    """,
                    parameters,
                )
                while rows := cursor.fetchmany(1000):
                    _raise_if_cancelled(is_cancelled)
                    for row in rows:
                        _raise_if_cancelled(is_cancelled)
                        record_count += 1
                        status = _report_status(row["status"])
                        reasons = _parse_reason_codes(
                            row["decision_reasons_json"]
                        )
                        status_counts[status] += 1
                        reason_counts.update(reasons)
                        worksheet.append(
                            _detail_cells(
                                worksheet,
                                row,
                                status=status,
                                reasons=reasons,
                            )
                        )
                        if (
                            progress_callback is not None
                            and (
                                record_count == total_count
                                or record_count % 100 == 0
                            )
                        ):
                            progress_callback(record_count, total_count)
                connection.execute("COMMIT")
        except InspectionExcelReportError:
            raise
        except sqlite3.Error as exc:
            raise InspectionExcelReportError(
                f"Unable to read inspection history for export: {exc}"
            ) from exc

        worksheet.auto_filter.ref = f"A1:R{record_count + 1}"
        return record_count, status_counts, reason_counts

    def _connect(self) -> sqlite3.Connection:
        uri = f"{self.database_path.resolve().as_uri()}?mode=ro"
        connection = sqlite3.connect(
            uri,
            uri=True,
            timeout=5.0,
            check_same_thread=True,
        )
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA query_only=ON")
        connection.execute("PRAGMA busy_timeout=5000")
        return connection

    @staticmethod
    def _write_summary(
        worksheet,
        filters: InspectionHistoryFilters,
        record_count: int,
        status_counts: Counter[str],
        *,
        language: str,
    ) -> None:
        zh = language == "zh"
        worksheet.column_dimensions["A"].width = 24
        worksheet.column_dimensions["B"].width = 48
        title = "AI 檢測紀錄報表" if zh else "AI Inspection Report"
        worksheet.append([_title_cell(worksheet, title), ""])
        worksheet.append([])
        generated_at = datetime.now(HISTORY_TIMEZONE).strftime(
            "%Y-%m-%d %H:%M:%S %Z"
        )
        passed = status_counts["PASS"]
        failed = status_counts["NG"]
        errors = status_counts["ERROR"]
        judged = passed + failed
        yield_rate = None if judged == 0 else passed / judged
        rows = (
            ("產生時間" if zh else "Generated at", generated_at),
            (
                "產品" if zh else "Product",
                _excel_text(filters.product) or ("全部" if zh else "All"),
            ),
            (
                "工位" if zh else "Station",
                _excel_text(filters.station) or ("全部" if zh else "All"),
            ),
            (
                "狀態" if zh else "Status",
                _excel_text(filters.status) or ("全部" if zh else "All"),
            ),
            (
                "起始時間" if zh else "From",
                _format_boundary(filters.started_at),
            ),
            (
                "結束時間（不含）" if zh else "Through (exclusive)",
                _format_boundary(filters.ended_at),
            ),
            ("符合筆數" if zh else "Matching records", record_count),
            ("PASS", passed),
            ("NG", failed),
            ("錯誤" if zh else "Errors", errors),
            ("良率" if zh else "Yield", yield_rate if yield_rate is not None else "--"),
        )
        for label, value in rows:
            label_cell = WriteOnlyCell(worksheet, value=label)
            label_cell.font = Font(bold=True, color="334E68")
            value_cell = WriteOnlyCell(worksheet, value=value)
            if label in {"良率", "Yield"} and isinstance(value, float):
                value_cell.number_format = "0.00%"
            worksheet.append([label_cell, value_cell])

    @staticmethod
    def _write_reason_summary(
        worksheet,
        record_count: int,
        reason_counts: Counter[str],
        *,
        language: str,
    ) -> None:
        worksheet.column_dimensions["A"].width = 36
        worksheet.column_dimensions["B"].width = 14
        worksheet.column_dimensions["C"].width = 16
        columns = (
            ("異常原因", "次數", "占全部檢測")
            if language == "zh"
            else ("Failure reason", "Count", "Share of records")
        )
        worksheet.append(_header_cells(worksheet, columns))
        worksheet.freeze_panes = "A2"
        for reason, count in reason_counts.most_common():
            ratio = count / record_count if record_count else 0
            ratio_cell = WriteOnlyCell(worksheet, value=ratio)
            ratio_cell.number_format = "0.00%"
            worksheet.append([_excel_text(reason), count, ratio_cell])
        worksheet.auto_filter.ref = f"A1:C{len(reason_counts) + 1}"


def _header_cells(worksheet, values: tuple[str, ...]) -> list[WriteOnlyCell]:
    cells: list[WriteOnlyCell] = []
    for value in values:
        cell = WriteOnlyCell(worksheet, value=value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="245B8F")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cells.append(cell)
    return cells


def _title_cell(worksheet, value: str) -> WriteOnlyCell:
    cell = WriteOnlyCell(worksheet, value=value)
    cell.font = Font(bold=True, size=16, color="1F3347")
    return cell


def _detail_cells(
    worksheet,
    row: sqlite3.Row,
    *,
    status: str,
    reasons: tuple[str, ...],
) -> list[object]:
    inference_time = row["inference_time"]
    values: list[object] = [
        _excel_text(row["inspection_id"]),
        _excel_text(_format_timestamp(str(row["timestamp"] or ""))),
        status,
        _excel_text(row["product"]),
        _excel_text(row["station"]),
        _excel_text(row["detector"]),
        (
            round(float(inference_time) * 1000.0, 3)
            if inference_time is not None
            else None
        ),
        _excel_text(row["machine_id"]),
        _excel_text(row["work_order"]),
        _excel_text(row["camera_id"]),
        _excel_text(row["model_version"]),
        _excel_text("；".join(reasons)),
        _excel_text(row["review_outcome"]),
        _excel_text(row["failure_category"]),
    ]
    cells: list[object] = list(values)
    for column_name in (
        "original_path",
        "annotated_path",
        "heatmap_path",
        "snapshot_path",
    ):
        raw_path = str(row[column_name] or "")
        cell = WriteOnlyCell(worksheet, value=_excel_text(raw_path))
        if raw_path:
            try:
                cell.hyperlink = Path(raw_path).resolve().as_uri()
                cell.style = "Hyperlink"
            except ValueError:
                pass
        cells.append(cell)
    return cells


def _parse_reason_codes(raw: object) -> tuple[str, ...]:
    try:
        payload = json.loads(str(raw or "[]"))
    except (TypeError, ValueError, json.JSONDecodeError):
        fallback = str(raw or "").strip()
        return (fallback,) if fallback else ()
    if not isinstance(payload, list):
        return (str(payload),)
    reasons: list[str] = []
    for item in payload:
        if isinstance(item, dict):
            value = item.get("code") or item.get("reason") or item.get("message")
            if value:
                reasons.append(str(value).strip())
        elif str(item).strip():
            reasons.append(str(item).strip())
    return tuple(dict.fromkeys(value for value in reasons if value))


def _report_status(value: object) -> str:
    normalized = str(value or "").strip().upper()
    if normalized in {"FAIL", "DETECTION_FAIL", "NG"}:
        return "NG"
    if normalized in {"ERROR", "INFERENCE_ERROR"}:
        return "ERROR"
    return normalized or "UNKNOWN"


def _format_timestamp(value: str) -> str:
    parsed = parse_inspection_timestamp(value)
    if parsed is None:
        return value
    return parsed.strftime("%Y-%m-%d %H:%M:%S")


def _format_boundary(value: datetime | None) -> str:
    if value is None:
        return "--"
    return value.astimezone(HISTORY_TIMEZONE).strftime("%Y-%m-%d %H:%M:%S")


def _set_detail_widths(worksheet) -> None:
    widths = (38, 21, 11, 18, 12, 14, 17, 14, 16, 14, 18, 34, 18, 20, 48, 48, 48, 48)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            chr(64 + index)
        ].width = width


def _excel_text(value: object) -> str:
    """Prevent database text from being interpreted as an Excel formula."""
    text = str(value or "")
    return f"'{text}" if text.startswith(("=", "+", "-", "@")) else text


def _raise_if_cancelled(
    is_cancelled: Callable[[], bool] | None,
) -> None:
    if is_cancelled is not None and is_cancelled():
        raise InspectionExcelReportCancelled(
            "Inspection report export was cancelled."
        )


def _discard_workbook(
    workbook: Workbook | None,
    temporary: Path,
) -> None:
    """Close write-only XML streams before removing a cancelled report."""
    if workbook is not None:
        try:
            workbook.save(temporary)
        except (OSError, RuntimeError, TypeError, ValueError):
            workbook.close()
    temporary.unlink(missing_ok=True)
